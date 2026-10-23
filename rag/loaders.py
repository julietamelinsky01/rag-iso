"""Lectura de documentos.

load_document() devuelve (metadatos, bloques). Cada bloque es un título,
un párrafo o una tabla, en el mismo orden en que aparecen en el documento.

- DOCX: títulos (estilos de Word y líneas cortas en negrita o subrayado),
  párrafos, viñetas y TABLAS (fila por fila, con el nombre de cada
  columna). Del encabezado de página toma código, revisión, fecha,
  emisión, quién lo realizó y quién lo aprobó.
  Word no guarda números de página fijos: en DOCX la página queda vacía.
- PDF: texto por página. Las tablas pueden perder su estructura.
- TXT / MD: texto plano, sin páginas; en MD, las líneas con # son títulos.
"""
import re
from dataclasses import asdict, dataclass
from pathlib import Path


@dataclass
class Block:
    kind: str                  # "heading" | "text" | "table"
    text: str
    level: int = 0             # nivel del título (1 = principal)
    page: int | None = None    # solo se conoce en PDF


@dataclass
class DocMeta:
    titulo: str | None = None
    codigo: str | None = None
    revision: str | None = None
    fecha_revision: str | None = None
    emision: str | None = None
    realizado: str | None = None
    aprobado: str | None = None

    def as_dict(self) -> dict:
        return asdict(self)

    def label(self) -> str:
        """Ej.: 'PRO-02 Liderazgo (Rev. 07, 12/11/2020)'."""
        label = " ".join(p for p in (self.codigo, self.titulo) if p)
        if self.revision:
            label += f" (Rev. {self.revision}"
            label += f", {self.fecha_revision})" if self.fecha_revision else ")"
        return label


def load_document(path: str) -> tuple[DocMeta, list[Block]]:
    ext = Path(path).suffix.lower()
    if ext == ".docx":
        return _load_docx(path)
    if ext == ".pdf":
        return DocMeta(titulo=Path(path).stem), _load_pdf(path)
    if ext in (".txt", ".md"):
        return DocMeta(titulo=Path(path).stem), _load_text(path)
    raise ValueError(
        f"Formato no soportado todavía: {ext} (usá PDF, DOCX, TXT o MD)"
    )


# ---------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------

PSEUDO_HEADING_LEVEL = 9  # títulos "a mano": debajo de los de estilo

# Word viejo mete palabras dentro de "smart tags", hipervínculos o campos
# que p.text / p.runs de python-docx no ven. Por eso se lee el XML.
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _load_docx(path: str) -> tuple[DocMeta, list[Block]]:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(path)
    meta = _docx_meta(doc, path)
    blocks: list[Block] = []

    for element in doc.element.body.iterchildren():
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, doc)
            text = _para_text(paragraph)
            if not text:
                continue
            level = _heading_level(paragraph)
            if level:
                blocks.append(Block("heading", text.rstrip(" .:"), level))
            elif element.find(".//" + qn("w:numPr")) is not None:
                blocks.append(Block("text", f"- {text}"))  # viñeta o lista
            else:
                blocks.append(Block("text", text))
        elif element.tag == qn("w:tbl"):
            table_text = _table_to_text(Table(element, doc))
            if table_text:
                blocks.append(Block("table", table_text))

    if not blocks:
        raise ValueError("El documento Word no tiene texto.")
    return meta, blocks


def _heading_level(paragraph) -> int:
    style = (
        (paragraph.style.name or "").lower()
        if paragraph.style is not None else ""
    )
    match = re.match(r"(heading|título|titulo)\s*(\d)", style)
    if match:
        return int(match.group(2))
    if style in ("title", "título", "titulo"):
        return 1
    # Título "a mano": línea corta con todo en negrita o todo subrayado.
    text = _para_text(paragraph)
    runs = _text_runs(paragraph)
    if runs and len(text) <= 120 and not text.endswith(":"):
        if (all(_run_flag(r, "b") for r in runs)
                or all(_run_flag(r, "u") for r in runs)):
            return PSEUDO_HEADING_LEVEL
    return 0


def _para_text(paragraph) -> str:
    parts = []
    for node in paragraph._p.iter():
        if node.tag == _W + "t" and node.text:
            parts.append(node.text)
        elif node.tag in (_W + "tab", _W + "br", _W + "cr"):
            parts.append(" ")
    return _norm("".join(parts))


def _text_runs(paragraph) -> list:
    return [
        run for run in paragraph._p.iter(_W + "r")
        if "".join(t.text or "" for t in run.iter(_W + "t")).strip()
    ]


def _run_flag(run, flag: str) -> bool:
    """True si la corrida tiene negrita (b) o subrayado (u) activo."""
    properties = run.find(_W + "rPr")
    if properties is None:
        return False
    element = properties.find(_W + flag)
    if element is None:
        return False
    return element.get(_W + "val") not in ("0", "false", "none")


def _cell_text(cell) -> str:
    parts = [_para_text(p) for p in cell.paragraphs]
    return " / ".join(part for part in parts if part)


def _is_header_cell(cell) -> bool:
    paragraphs = [p for p in cell.paragraphs if _para_text(p)]
    if not paragraphs:
        return True  # celda vacía de una fila de títulos (celdas combinadas)
    for paragraph in paragraphs:
        style = (
            (paragraph.style.name or "").lower()
            if paragraph.style is not None else ""
        )
        runs = _text_runs(paragraph)
        if not (
            style.startswith(("heading", "título", "titulo"))
            or (runs and all(_run_flag(r, "b") for r in runs))
        ):
            return False
    return True


def _unique_cells(row):
    """Las celdas combinadas aparecen repetidas: se marca la repetición
    pero se conserva la posición para alinear con los títulos."""
    seen, cells = set(), []
    for cell in row.cells:
        key = id(cell._tc)
        cells.append((cell, key in seen))
        seen.add(key)
    return cells


def _table_to_text(table) -> str:
    rows = list(table.rows)
    if not rows:
        return ""
    ncols = max(len(r.cells) for r in rows)

    # Tabla clave | valor: 2 columnas sin fila de títulos.
    if ncols == 2 and not all(_is_header_cell(c) for c in rows[0].cells[1:]):
        lines = []
        for row in rows:
            key, value = _cell_text(row.cells[0]), _cell_text(row.cells[1])
            if key and value and key != value:
                lines.append(f"{key}: {value}")
            elif key or value:
                lines.append(key or value)
        return "\n".join(lines)

    # Filas de títulos al principio (pueden ser 2 por celdas combinadas).
    n_header = 0
    for row in rows[:3]:
        if (all(_is_header_cell(c) for c in row.cells)
                and any(c.text.strip() for c in row.cells)):
            n_header += 1
        else:
            break
    n_header = n_header or 1  # sin formato de título: la primera fila

    headers = [""] * ncols
    for row in rows[:n_header]:
        for j, (cell, _) in enumerate(_unique_cells(row)):
            text = _cell_text(cell)
            if text and text not in headers[j].split(" - "):
                headers[j] = f"{headers[j]} - {text}" if headers[j] else text

    lines = []
    for row in rows[n_header:]:
        pairs, last = [], None
        for j, (cell, repeated) in enumerate(_unique_cells(row)):
            text = _cell_text(cell)
            if not text or repeated or text == last:
                continue
            last = text
            header = headers[j] if j < len(headers) else ""
            pairs.append(f"{header}: {text}" if header else text)
        if pairs:
            # Un dato por renglón: el modelo relaciona mejor cada columna
            # con la pregunta que en una fila larga separada por ";".
            lines.append("\n".join([pairs[0]] + [f"  - {p}" for p in pairs[1:]]))
    if not lines:  # tabla sin datos de texto: se dejan los títulos
        lines.append(" | ".join(h for h in headers if h))
    # Una línea en blanco separa las filas.
    return "\n\n".join(lines)


_HEADER_LABELS = re.compile(
    r"c[óo]digo|emisi[óo]n|rev\.|revisi[óo]n|fecha|realizado|aprobado|p[áa]gina",
    re.IGNORECASE,
)


def _docx_meta(doc, path: str) -> DocMeta:
    texts = []
    for section in doc.sections[:1]:
        header = section.header
        texts += [_para_text(p) for p in header.paragraphs]
        for table in header.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = _cell_text(cell)
                    if text not in texts:  # celdas combinadas se repiten
                        texts.append(text)
    flat = re.sub(r"\s+", " ", "\n".join(texts))

    def find(pattern: str) -> str | None:
        match = re.search(pattern, flat, re.IGNORECASE)
        return match.group(1).strip() if match else None

    codigo = find(r"C[óo]digo\s*:?\s*([A-Z]{2,5}\s*-?\s*\d{1,4})")
    if codigo:
        codigo = re.sub(r"\s*-?\s*(\d+)$", r"-\1", re.sub(r"\s+", "", codigo))

    meta = DocMeta(
        codigo=codigo,
        revision=find(r"Rev(?:isi[óo]n)?\.?\s*:?\s*(\d{1,3})\b"),
        fecha_revision=find(
            r"Rev(?:isi[óo]n)?\.?\s*:?\s*\d{1,3}[\s/]+Fecha\s*:?\s*"
            r"(\d{1,2}/\d{1,2}/\d{2,4})"
        ),
        emision=find(r"Emisi[óo]n\s*:?\s*(\d{1,2}/\d{1,2}/\d{2,4})"),
        realizado=find(
            r"Realizado\s+por\s*:?\s*(.+?)\s*(?=Aprobado\s+por|P[áa]gina|$)"
        ),
        aprobado=find(
            r"Aprobado\s+por\s*:?\s*(.+?)\s*(?=Realizado\s+por|P[áa]gina|$)"
        ),
    )
    for text in texts:
        text = _norm(text)
        if text and not _HEADER_LABELS.search(text) and len(text) <= 100:
            meta.titulo = text
            break
    meta.titulo = meta.titulo or Path(path).stem
    return meta


# ---------------------------------------------------------------
# PDF / TXT / MD
# ---------------------------------------------------------------

def _load_pdf(path: str) -> list[Block]:
    from pypdf import PdfReader

    blocks = []
    for number, page in enumerate(PdfReader(path).pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            blocks.append(Block("text", text, page=number))
    if not blocks:
        raise ValueError(
            "El PDF no tiene texto extraíble (¿es escaneado?). Hace falta OCR, "
            "por ejemplo con Docling."
        )
    return blocks


def _load_text(path: str) -> list[Block]:
    blocks = []
    content = Path(path).read_text(encoding="utf-8", errors="ignore")
    for line in content.splitlines():
        match = re.match(r"^(#{1,6})\s+(.*)", line)
        if match:
            blocks.append(
                Block("heading", match.group(2).strip(), len(match.group(1)))
            )
        elif line.strip():
            blocks.append(Block("text", line))
    return blocks


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()
