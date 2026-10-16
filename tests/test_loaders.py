"""Pruebas de lectura y fragmentación (rag/loaders.py y rag/chunker.py).

Crean un Word de ejemplo con la estructura de un procedimiento real
(encabezado, títulos en negrita, tabla e historial). No usan red.
Ejecutar desde la raíz del proyecto:
    python -m unittest discover -s tests -v
"""
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402

from rag.chunker import Block, chunk_blocks  # noqa: E402
from ingest import read_file  # noqa: E402
from rag.loaders import load_document  # noqa: E402


def bold(document, text):
    document.add_paragraph().add_run(text).bold = True


def build_procedure(path: Path) -> None:
    document = Document()
    header = document.sections[0].header
    table = header.add_table(rows=2, cols=3, width=document.sections[0].page_width)
    table.cell(0, 0).text = "Liderazgo"
    table.cell(0, 1).text = "Código PRO - 02"
    table.cell(0, 2).text = "Emisión 24/07/2008"
    table.cell(1, 0).text = "Rev. 07 Fecha 12/11/2020"
    table.cell(1, 1).text = "Realizado por: Responsable de Calidad"
    table.cell(1, 2).text = "Aprobado por: Dirección"

    bold(document, "Descripción del Proceso")
    bold(document, "Revisión por la Dirección")
    document.add_paragraph(
        "Las reuniones de revisión se realizan al menos una vez al año."
    )
    bold(document, "Objetivos")
    document.add_paragraph("La revisión de objetivos se realiza semestralmente.")

    bold(document, "Registros")
    records = document.add_table(rows=3, cols=3)
    for j, title in enumerate(("Identificación", "Soporte", "Tiempo")):
        records.cell(0, j).paragraphs[0].add_run(title).bold = True
    for i, row in enumerate(
        (("Matriz FODA", "Electrónico", "Permanente"),
         ("Minuta de Reunión", "Papel", "5 años")),
        start=1,
    ):
        for j, value in enumerate(row):
            records.cell(i, j).text = value

    bold(document, "Historial de Cambios")
    history = document.add_table(rows=2, cols=2)
    history.cell(0, 0).paragraphs[0].add_run("Versión").bold = True
    history.cell(0, 1).paragraphs[0].add_run("Descripción").bold = True
    history.cell(1, 0).text = "07"
    history.cell(1, 1).text = "La revisión de objetivos pasa a ser semestral."
    document.save(path)


class DocxTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = path = Path(cls.tmp.name) / "PRO-02.docx"
        build_procedure(path)
        cls.meta, blocks = load_document(str(path))
        cls.chunks = chunk_blocks(blocks)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def find(self, text):
        return next(c for c in self.chunks if text in c.content)

    def test_encabezado(self):
        self.assertEqual(self.meta.codigo, "PRO-02")
        self.assertEqual(self.meta.revision, "07")
        self.assertEqual(self.meta.fecha_revision, "12/11/2020")
        self.assertEqual(self.meta.emision, "24/07/2008")
        self.assertEqual(self.meta.realizado, "Responsable de Calidad")
        self.assertEqual(self.meta.aprobado, "Dirección")
        self.assertEqual(self.meta.label(), "PRO-02 Liderazgo (Rev. 07, 12/11/2020)")

    def test_frecuencias_en_secciones_distintas(self):
        annual = self.find("al menos una vez al año")
        semester = self.find("semestralmente")
        self.assertIsNot(annual, semester)
        self.assertTrue(annual.titulo.endswith("Revisión por la Dirección"))
        self.assertTrue(semester.titulo.endswith("Objetivos"))

    def test_tabla_fila_por_fila_con_columnas(self):
        chunk = self.find("Matriz FODA")
        self.assertIn(
            "Identificación: Matriz FODA; Soporte: Electrónico; Tiempo: Permanente",
            chunk.content,
        )
        self.assertEqual(chunk.titulo, "Registros")

    def test_historial_marcado(self):
        chunk = self.find("pasa a ser semestral")
        self.assertEqual(chunk.tipo, "historial")

    def test_ficha_del_encabezado_se_puede_buscar(self):
        _, chunks = read_file(self.path)
        self.assertEqual(chunks[0].titulo, "Ficha del documento")
        self.assertIn("Aprobado por: Dirección", chunks[0].content)
        self.assertIn("Revisión: 07 del 12/11/2020", chunks[0].content)

    def test_word_sin_paginas(self):
        self.assertTrue(all(c.pagina is None for c in self.chunks))


class PdfTests(unittest.TestCase):
    def test_clausulas_con_pagina(self):
        chunks = chunk_blocks([
            Block("text", "3.3 Auditorías internas\nSe realizan cada cuatro meses.", page=3),
        ])
        self.assertEqual(chunks[0].clausula, "3.3")
        self.assertEqual(chunks[0].pagina, 3)

    def test_pagina_de_indice_se_omite(self):
        toc = "\n".join(
            f"{n}.1 Título de la sección número {n} {n + 10}" for n in range(1, 8)
        )
        chunks = chunk_blocks([
            Block("text", f"Índice\n{toc}", page=2),
            Block("text", "1.1 Título de la sección número 1\nTexto real.", page=11),
        ])
        self.assertEqual([(c.clausula, c.pagina) for c in chunks], [("1.1", 11)])


if __name__ == "__main__":
    unittest.main()
